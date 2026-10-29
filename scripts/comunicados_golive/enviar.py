"""Comunicados de go-live de cliente (29/10/2026; 1º uso: VIXPAR — go-live 04/11).

  python3 scripts/comunicados_golive/enviar.py --teste voce@empresa.com.br        # 1 teste de cada comunicado
  python3 scripts/comunicados_golive/enviar.py --importar VIX planilha.xlsx         # carrega/atualiza a lista
  python3 scripts/comunicados_golive/enviar.py --excluir VIX 12345678000199 ... --motivo "participou dos testes"
  python3 scripts/comunicados_golive/enviar.py --resumo VIX                         # quem recebe / quem está fora
  python3 scripts/comunicados_golive/enviar.py --validar VIX                        # lista conferida com o cliente (trava)
  python3 scripts/comunicados_golive/enviar.py --enviar VIX 1 [--dry-run]           # comunicado N para a lista
  python3 scripts/comunicados_golive/enviar.py --agendado                           # comunicado do CALENDARIO de hoje

Destinatários: golive_contacts (patch_162). Excluídos (ex.: participantes dos
testes) nunca recebem. Cada envio fica em envios->'N' — reexecutar não reenvia.
Remetente SIGEC-ELOS; responder vai para o suporte. Nada de nome real de
cliente fora de CAMPANHAS (os textos são da campanha).
"""
import argparse, datetime, json, os, re, shutil, subprocess, sys, time, urllib.request, urllib.error
from zoneinfo import ZoneInfo

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '../..'))
_envf = os.path.join(ROOT, '.env')
ENV = {l.split('=', 1)[0]: l.split('=', 1)[1].strip().strip('"') for l in open(_envf) if '=' in l and not l.startswith('#')} if os.path.exists(_envf) else {}
DB = os.environ.get('MIGRACAO_DB_URL') or ENV.get('SUPABASE_DB_URL')
RESEND = os.environ.get('RESEND_API_KEY') or ENV.get('RESEND_API_KEY')
PSQL = shutil.which('psql') or '/usr/local/opt/libpq/bin/psql'
FROM = 'SIGEC-ELOS <noreply@eqpitech.com.br>'
REPLY_TO = 'suporte@eqpitech.com.br'
SITE = 'https://elos.eqpitech.com.br'
SUPORTE = ('<strong>E-mail:</strong> suporte@eqpitech.com.br · <strong>WhatsApp:</strong> (11) 5039-6769 · '
           '<strong>Telefone:</strong> (11) 2305-3881 (segunda a sexta, das 8h às 18h)')

# Campanhas de go-live. 'custeio': quem paga a homologação (texto do e-mail).
CAMPANHAS = {
    'VIX': {'cliente': 'VIXPAR', 'inicio': 'quarta-feira, 04/11', 'inicio_curto': '04/11',
            'portal': f'{SITE}/portal/vixpar', 'custeio': 'cliente'},
}
# data em Brasília → (campanha, comunicado). Workflow comunicados-golive.yml (--agendado)
CALENDARIO = {'2026-10-30': ('VIX', 1), '2026-11-03': ('VIX', 2)}
MANUAL = f'{SITE}/manual/fornecedor'

def psql(sql):
    r = subprocess.run([PSQL, DB, '-X', '-q', '-At', '-v', 'ON_ERROR_STOP=1', '-c', sql],
                       capture_output=True, text=True, env={**os.environ, 'PGGSSENCMODE': 'disable'})
    if r.returncode: raise SystemExit(r.stderr)
    return r.stdout

def q(v): return "'" + str(v).replace("'", "''") + "'" if v is not None else 'null'

def moldura(k, titulo, corpo):
    return f"""<div style="font-family:Arial,Helvetica,sans-serif;max-width:600px;margin:0 auto;color:#374151">
  <div style="background:#2E3192;padding:26px 24px;border-radius:12px 12px 0 0;text-align:center">
    <div style="color:#fff;font-size:22px;font-weight:bold;letter-spacing:.5px">SIGEC-ELOS</div>
    <div style="color:#F47E2F;font-size:13px;margin-top:4px">{titulo}</div></div>
  <div style="background:#fff;padding:28px 26px;border:1px solid #e2e8f0;border-top:none;font-size:15px;line-height:1.65">{corpo}
    <p style="margin-top:22px">Conte com a gente,<br><strong>Equipe SIGEC-ELOS · EQPI Tech</strong></p></div>
  <div style="background:#f8fafc;padding:16px;border:1px solid #e2e8f0;border-top:none;border-radius:0 0 12px 12px;text-align:center;font-size:12px;color:#6b7280;line-height:1.6">
    Dúvidas? Fale com o suporte — {SUPORTE}<br>EQPI Tech · elos.eqpitech.com.br</div></div>"""

def botao(txt, href):
    return f'<p style="text-align:center;margin:24px 0"><a href="{href}" style="display:inline-block;background:#F47E2F;color:#fff;padding:14px 34px;border-radius:9px;text-decoration:none;font-weight:bold;font-size:15px">{txt}</a></p>'

def lista(itens, ordenada=False):
    tag = 'ol' if ordenada else 'ul'
    return f'<{tag} style="padding-left:20px;margin:6px 0 16px">' + ''.join(f'<li style="margin-bottom:6px">{i}</li>' for i in itens) + f'</{tag}>'

def ola(c): return f"Olá, <strong>{c['razao_social']}</strong>!" if c.get('razao_social') and (c.get('empresas') or 1) == 1 else 'Olá!'

# patch_169: o texto depende de como a pessoa entra em 04/11
def modo(c): return c.get('modo') or 'convite'
def suas(c):
    n = c.get('empresas') or 1
    return f'as suas <strong>{n} empresas</strong>' if n > 1 else 'a sua empresa'

def custo(k):
    return ('<p style="background:#f0fdf4;border-left:4px solid #22c55e;padding:10px 14px;border-radius:0 8px 8px 0;margin:16px 0">'
            f'<strong>Sem custo para a sua empresa:</strong> a homologação é custeada pela {k["cliente"]}.</p>') if k['custeio'] == 'cliente' else ''

# ── Comunicado 1 (30/10): o aviso, com o manual ─────────────────────────────
def passos1(k, c):
    if modo(c) == 'precadastro':
        return [f'Para facilitar, a EQPI deixa {suas(c)} <strong>já cadastrada{"s" if (c.get("empresas") or 1) > 1 else ""}</strong> no SIGEC-ELOS, com o processo da {k["cliente"]} aberto;',
                f'Em <strong>{k["inicio_curto"]}</strong>, você recebe neste e-mail o <strong>seu acesso</strong> — é só criar a senha' + (' e escolher a empresa no menu' if (c.get('empresas') or 1) > 1 else '') + ';',
                'Na véspera, enviamos um lembrete.']
    if modo(c) == 'operacional':
        return [f'Em <strong>{k["inicio_curto"]}</strong>, o responsável pela sua empresa recebe o <strong>convite da {k["cliente"]}</strong>;',
                'Assim que ele concluir o cadastro, você recebe neste e-mail o <strong>seu acesso</strong> à plataforma;',
                'Na véspera, enviamos um lembrete.']
    return [f'Em <strong>{k["inicio_curto"]}</strong>, você recebe neste e-mail o <strong>convite da {k["cliente"]}</strong>, com o link para criar o seu acesso e enviar os documentos;',
            'Na véspera, enviamos um lembrete.']

def com1(k, c):
    assunto = f"{k['cliente']}: a partir de {k['inicio_curto']}, a homologação de fornecedores é pelo SIGEC-ELOS"
    corpo = f"""<p>{ola(c)}</p>
<p>A partir de <strong>{k['inicio']}</strong>, a homologação dos fornecedores da <strong>{k['cliente']}</strong> passa a ser feita no
<strong>SIGEC-ELOS</strong>, a plataforma online da EQPI Tech. Estar homologado no SIGEC-ELOS passa a ser requisito para fornecer à {k['cliente']}.</p>
<p style="margin-bottom:4px"><strong>O que melhora para a sua empresa</strong></p>
{lista(['Cadastro e documentos em um só lugar, online;',
        'Acompanhamento do processo em tempo real;',
        'Aviso antes do vencimento de cada documento;',
        'Várias certidões (CNPJ, FGTS, trabalhista, federal e outras) são obtidas automaticamente nas fontes oficiais — você não precisa enviar.'])}
<p style="margin-bottom:4px"><strong>Os próximos passos</strong></p>
{lista(passos1(k, c), ordenada=True)}
{custo(k)}
<p>Para conhecer a plataforma antes, veja o <strong>manual do fornecedor</strong>, com o passo a passo e as imagens das telas:</p>
{botao('Abrir o manual do fornecedor', MANUAL)}
<p><strong>Não é preciso fazer nada agora</strong> — aguarde o convite.</p>"""
    return assunto, moldura(k, f"Homologação de fornecedores {k['cliente']} · Comunicado 1 de 2", corpo)

# ── Comunicado 2 (03/11): amanhã chega o convite ────────────────────────────
def com2(k, c):
    m = modo(c)
    o_que = {'precadastro': f'o <strong>seu acesso ao SIGEC-ELOS</strong>, com {suas(c)} já cadastrada{"s" if (c.get("empresas") or 1) > 1 else ""}',
             'operacional': f'o responsável pela sua empresa recebe o <strong>convite da {k["cliente"]}</strong>; quando ele concluir o cadastro, você recebe o seu acesso',
             'convite': f'você recebe neste e-mail o <strong>convite da {k["cliente"]}</strong> para a homologação no SIGEC-ELOS'}[m]
    assunto = (f"Amanhã ({k['inicio_curto']}): o seu acesso ao SIGEC-ELOS para a homologação com a {k['cliente']}" if m == 'precadastro'
               else f"Amanhã ({k['inicio_curto']}): o convite da {k['cliente']} para a sua homologação no SIGEC-ELOS")
    primeiro = ('<strong>Criar a sua senha</strong> pelo link do e-mail de acesso' + (' e escolher a empresa no menu' if (c.get('empresas') or 1) > 1 else '') + ';') if m == 'precadastro' \
        else '<strong>Criar a sua senha</strong> pelo link do e-mail de acesso, que chega depois que o responsável concluir o cadastro;' if m == 'operacional' \
        else '<strong>Criar o seu acesso</strong> — com o CNPJ e o seu e-mail, em poucos minutos;'
    corpo = f"""<p>{ola(c)}</p>
<p>Como avisamos, <strong>amanhã, {k['inicio']}</strong>, {'chega neste e-mail ' if m == 'precadastro' else ''}{o_que}.</p>
<p style="margin-bottom:4px"><strong>{'No SIGEC-ELOS, você vai:' if m != 'convite' else 'Pelo link do convite, você vai:'}</strong></p>
{lista([primeiro,
        '<strong>Ativar a verificação em duas etapas</strong> — um código no celular, para a sua segurança;',
        '<strong>Enviar os documentos</strong> pedidos para a sua atividade e responder o questionário;',
        '<strong>Acompanhar a análise</strong> — avisamos a cada etapa.'], ordenada=True)}
{custo(k)}
<p><strong>Para adiantar:</strong> deixe à mão o contrato social e os documentos da empresa.</p>
<p style="font-size:14px">{'O seu acesso chega por e-mail assim que o responsável concluir o cadastro da empresa. Não chegou? Confira a caixa de spam ou fale com o suporte.' if m == 'operacional' else f"Não recebeu {'o e-mail de acesso' if m == 'precadastro' else 'o convite'} amanhã? Confira a caixa de spam ou fale com o suporte."} Dúvidas no passo a passo:</p>
{botao('Abrir o manual do fornecedor', MANUAL)}"""
    return assunto, moldura(k, f"Homologação de fornecedores {k['cliente']} · Comunicado 2 de 2", corpo)

COMUNICADOS = {1: com1, 2: com2}

def resend(to, assunto, html):
    req = urllib.request.Request('https://api.resend.com/emails', method='POST', data=json.dumps({
        'from': FROM, 'reply_to': REPLY_TO, 'to': [to], 'subject': assunto, 'html': html}).encode(),
        headers={'Content-Type': 'application/json', 'Authorization': f'Bearer {RESEND}', 'User-Agent': 'sigec-elos-golive/1.0'})
    with urllib.request.urlopen(req, timeout=30) as r: return json.loads(r.read()).get('id')

def contatos(camp):
    out = psql(f"select coalesce(json_agg(row_to_json(c)),'[]') from (select id, cnpj, razao_social, email, envios, modo, empresas from golive_contacts "
               f"where campanha={q(camp)} and not excluido order by razao_social) c")
    return json.loads(out or '[]')

def validada(camp):
    return psql(f"select coalesce((value->>'validada')::boolean, false) from app_settings where key={q('golive_lista_' + camp)}").strip() == 't'

def manual_no_ar():
    class SemRedirect(urllib.request.HTTPRedirectHandler):
        def redirect_request(self, *a, **k): return None
    try:
        urllib.request.build_opener(SemRedirect).open(urllib.request.Request(MANUAL, headers={'User-Agent': 'sigec-elos-golive/1.0'}), timeout=20)
        return False                                    # 200 = SPA devolvendo index (rota não existe)
    except urllib.error.HTTPError as e:
        return e.code == 302

def enviar(camp, n, dry=False):
    k = CAMPANHAS[camp]
    # travas: lista validada pelo cliente (--validar) e manual público no ar
    if not dry and not validada(camp):
        raise SystemExit(f'{camp}: lista NÃO validada — rode --validar {camp} depois da planilha final do cliente. Nada enviado.')
    if not dry and not manual_no_ar():
        raise SystemExit(f'{MANUAL} fora do ar — nada enviado.')
    feitos = falhas = pulados = 0
    for c in contatos(camp):
        if str(n) in (c.get('envios') or {}): pulados += 1; continue
        assunto, html = COMUNICADOS[n](k, c)
        if dry: feitos += 1; continue
        try:
            rid = resend(c['email'], assunto, html)
            psql(f"update golive_contacts set envios = envios || jsonb_build_object('{n}', jsonb_build_object('em', now(), 'id', {q(rid)})) where id={q(c['id'])}")
            feitos += 1
        except Exception as e:
            falhas += 1; print('falha', c['email'], e)
        time.sleep(0.6)                                  # limite da Resend
    print(f'{camp} comunicado {n}: {"simulados" if dry else "enviados"} {feitos} · já tinham recebido {pulados} · falhas {falhas}')

def importar(camp, arq):
    sys.path.insert(0, '/private/tmp')                  # openpyxl fora do repo, se preciso
    import openpyxl
    ws = openpyxl.load_workbook(arq, data_only=True).active
    cab = [str(c.value or '').strip().lower() for c in ws[1]]
    def col(*nomes):
        for i, h in enumerate(cab):
            if any(n in h for n in nomes): return i
        return None
    iC, iR, iE, iN, iX = col('cnpj'), col('razão', 'razao'), col('e-mail', 'email'), col('nível', 'nivel'), col('excluir', 'teste')
    n = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        email = str(row[iE] or '').strip().lower() if iE is not None else ''
        if not re.match(r'^[^@\s]+@[^@\s]+\.[^@\s]+$', email): continue
        cnpj = re.sub(r'\D', '', str(row[iC] or '')).zfill(14) if iC is not None and row[iC] else None
        excl = iX is not None and str(row[iX] or '').strip().lower() in ('sim', 's', 'x', 'true', '1')
        psql(f"insert into golive_contacts (campanha, cnpj, razao_social, email, nivel, excluido, motivo_exclusao) values "
             f"({q(camp)}, {q(cnpj)}, {q(row[iR] if iR is not None else None)}, {q(email)}, {q(row[iN] if iN is not None else None)}, "
             f"{'true' if excl else 'false'}, {q('marcado na planilha' if excl else None)}) "
             f"on conflict (campanha, email) do update set cnpj=excluded.cnpj, razao_social=excluded.razao_social, nivel=excluded.nivel, "
             f"excluido=golive_contacts.excluido or excluded.excluido, motivo_exclusao=coalesce(golive_contacts.motivo_exclusao, excluded.motivo_exclusao)")
        n += 1
    print(f'{camp}: {n} contatos carregados/atualizados')

if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--teste'); ap.add_argument('--variante'); ap.add_argument('--importar', nargs=2); ap.add_argument('--excluir', nargs='+')
    ap.add_argument('--motivo', default='participou dos testes'); ap.add_argument('--resumo')
    ap.add_argument('--validar'); ap.add_argument('--enviar', nargs=2); ap.add_argument('--dry-run', action='store_true'); ap.add_argument('--agendado', action='store_true')
    a = ap.parse_args()
    if a.teste:
        k = CAMPANHAS['VIX']
        variantes = [('convite', {'razao_social': 'EMPRESA EXEMPLO LTDA', 'modo': 'convite', 'empresas': 1}),
                     ('operacional', {'razao_social': 'EMPRESA EXEMPLO LTDA', 'modo': 'operacional', 'empresas': 1}),
                     ('pré-cadastro, 107 empresas', {'razao_social': None, 'modo': 'precadastro', 'empresas': 107})]
        so = [v for v in variantes if not a.variante or a.variante in v[0]]
        for nome, c in so:
            for n, f in COMUNICADOS.items():
                assunto, html = f(k, c)
                print('teste', nome, n, resend(a.teste, f'[TESTE {nome} — comunicado {n}, {"30/10" if n == 1 else "03/11"}] {assunto}', html))
    elif a.importar: importar(*a.importar)
    elif a.excluir:
        camp, cnpjs = a.excluir[0], [re.sub(r'\D', '', x).zfill(14) for x in a.excluir[1:]]
        print(psql(f"update golive_contacts set excluido=true, motivo_exclusao={q(a.motivo)} where campanha={q(camp)} and cnpj = any(array[{','.join(q(x) for x in cnpjs)}]) returning email").strip() or 'nenhum')
    elif a.resumo:
        print(psql(f"select case when excluido then 'FORA: '||coalesce(motivo_exclusao,'') else 'recebe' end, count(*) from golive_contacts where campanha={q(a.resumo)} group by 1"))
    elif a.validar:
        psql(f"insert into app_settings (key, value) values ({q('golive_lista_' + a.validar)}, jsonb_build_object('validada', true, 'em', now())) "
             f"on conflict (key) do update set value = excluded.value")
        print(a.validar, 'lista validada — os envios agendados podem sair')
    elif a.enviar: enviar(a.enviar[0], int(a.enviar[1]), a.dry_run)
    elif a.agendado:
        # o GitHub atrasa os agendamentos em até ~4h30 (visto em 14–16/10): o
        # workflow dispara de madrugada e aqui se espera até as 09:00 de Brasília.
        # Disparos de reserva (08:07 e 10:07 BRT) só completam quem faltou.
        agora = datetime.datetime.now(ZoneInfo('America/Sao_Paulo'))
        hoje = agora.date().isoformat()
        if hoje not in CALENDARIO: print('nada agendado para', hoje)
        else:
            alvo = agora.replace(hour=9, minute=0, second=0, microsecond=0)
            if agora < alvo:
                print(f'aguardando até {alvo:%H:%M} (Brasília)…', flush=True)
                time.sleep((alvo - agora).total_seconds())
            enviar(*CALENDARIO[hoje])
